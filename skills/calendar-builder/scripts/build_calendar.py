#!/usr/bin/env python3
"""Build an offline .xlsx copy of the content calendar from a rows JSON file.

Usage:
  python3 build_calendar.py rows.json out.xlsx [--config workspace.json] [--clean-titles] [--sheets-formulas]

The live calendar is always built natively with the connectors (see SKILL.md, Mode 1). This file is an
offline copy only: for review, for a backup, or as a starting point on a platform without a connector.
Never upload it as the live calendar.

rows.json (see examples/sample-rows.json):
  {"rows": [{"date": "YYYY-MM-DD", "ampm": "AM|PM|", "content_type", "channel", "title", "spec",
             "status", optional "week", "row_id", "notes", "missing", "sources", "visual", "proof_item"}],
   "milestones": [{"milestone", "type", "start", "end", "confirmed", "notes", "rows_proposed"}],
   "ideas": [{"idea", "added_by", "date_added", "wanted_by", "link", "status", "calendar_row", "agent_note"}]}

Week is counted from plan.week1_monday when a row has none. Row ID is W{week:02d}-{seq:02d} in date
order (AM before PM) when a row has none. Pillar and What to post come from the mappings in
references/calendar-schema.md, or from calendar.pillar_map and calendar.what_to_post_map in the config.
Statuses, proof types, proof placeholders and the team-owned pattern come from the config calendar
section, with the schema defaults when a key is missing.

--clean-titles     remove quote marks and put the proof placeholder on proof rows with no Proof item
                   (Mode 1 step 7: only after verify_calendar.py passes).
--sheets-formulas  write the Google Sheets formulas (Action needed, Draft, This Week) into the file,
                   for a person who imports it into Google Sheets by hand. Excel shows them as errors.
                   Without it, G, H and the This Week view are left empty.

Exit 0 on success, 2 on a usage or input error. Needs openpyxl.
"""
import json
import re
import sys
from datetime import date

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.stderr.write("build_calendar.py needs openpyxl: pip install openpyxl\n")
    sys.exit(2)

# Neutral defaults. Replace with the client's palette in their sheet, not in this repo.
HEADER, AGENT_HEADER = "1F2933", "6B7280"
STATUS_FILL = {"Proposed": ("FFF4CC", "000000"), "Draft ready": ("E3EEF7", "000000"),
               "Approved": ("1E6B52", "FFFFFF"), "Posted": ("E5E7EB", "000000"),
               "Cancelled": ("E5E7EB", "6B7280")}
OVERDUE, BAND, MISSING_TEXT = "F8D7D3", "F5F7F9", "9B1C1C"

DEFAULT_STATUSES = ["Proposed", "Planned", "Draft ready", "Approved", "Posted", "Cancelled"]
DEFAULT_TEAM_OWNED = "Publish|Newsletter|Distribute"
DEFAULT_PROOF_TYPES = {"Partner": "Partner Win", "Customer": "Customer Proof"}
DEFAULT_PLACEHOLDERS = {"Partner Win": "Partner news (to confirm)", "Customer Proof": "Customer story (to confirm)"}

COLS = ["Date", "Status", "What to post", "Pillar", "Channel", "Title / angle", "Action needed", "Draft",
        "Visual", "Notes", "What's missing", "Sources", "Row ID", "Week", "AM/PM", "Content type", "Spec",
        "Proof item", "Published link", "Final text", "Draft URL", "Original draft"]
WIDTHS = [14, 12, 30, 22, 14, 44, 30, 30, 30, 40, 30, 36, 9, 6, 7, 26, 50, 22, 22, 40, 30, 40]
WRAP = {"What to post", "Title / angle", "Action needed", "Visual", "Notes", "What's missing", "Sources"}

G2 = ('=ARRAYFORMULA(IF(A2:A="","",IF(B2:B="Posted","Done",IF(B2:B="Cancelled","Cancelled: no action",'
      'IF(A2:A<TODAY(),"Overdue: post it or update the status",IF(B2:B="Proposed","New slot: review, then set Planned",'
      'IF(REGEXMATCH(P2:P&"","<team>"),IF(A2:A=TODAY(),"Team: publish today, then set Posted",'
      '"Team writes this: set Posted when live"),IF(B2:B="Approved",IF(A2:A=TODAY(),"Post today",'
      '"Ready: post on the date"),IF(B2:B="Draft ready","Review the draft, then set Approved",'
      'IF(K2:K<>"","Waiting on: see What\'s missing",IF(A2:A-TODAY()<=7,"Draft due","")))))))))))')
DRAFT_LINK = ('=ARRAYFORMULA(IF({u}="","",HYPERLINK({u},IF({v}="","Open draft",IF(LEN(REGEXEXTRACT({v}&"","^[^\\n]*"))>80,'
              'LEFT(REGEXEXTRACT({v}&"","^[^\\n]*"),77)&"...",REGEXEXTRACT({v}&"","^[^\\n]*"))))))')
WEEK_FILTER = ('Calendar!A2:A<>"",Calendar!A2:A<=TODAY()+13,Calendar!B2:B<>"Cancelled",'
               '(Calendar!A2:A>=TODAY())+(Calendar!B2:B<>"Posted")')
TW_A3 = ('=IFERROR(CHOOSECOLS(SORT(FILTER({Calendar!A2:A,Calendar!B2:B,Calendar!C2:C,Calendar!E2:E,Calendar!F2:F,'
         'Calendar!G2:G,Calendar!K2:K},' + WEEK_FILTER + '),1,TRUE),1,2,3,4,5,6,7),"Nothing due in the next 14 days")')
TW_I3 = ('=IFERROR(CHOOSECOLS(SORT(FILTER({Calendar!I2:I,Calendar!J2:J,Calendar!U2:U,Calendar!V2:V,Calendar!A2:A},'
         + WEEK_FILTER + '),5,TRUE),1,2,3,4),"")')

HOW_TO_USE = [
    "This Week shows what is due in the next 14 days and anything overdue. Start there.",
    "Rows marked Proposed are new slots from the agents. Read the title and Notes, then set Planned to approve or Cancelled to drop.",
    "Agents draft Planned rows and set Draft ready. Click the Draft link to read the post.",
    "To change a Status, use the link at the top of This Week. It opens the Calendar showing only this week's posts.",
    "Set Approved when a draft is ready to post, and Posted once it is live.",
    "To drop a post, set its Status to Cancelled. Do not delete the row.",
    "Add post ideas on the Ideas tab. Set Declined on an idea instead of deleting it.",
    "Add trade shows, launches and other dates on the Milestones tab. To cancel a milestone, set Confirmed to Cancelled. Do not delete the row.",
    "Only change Status and Notes on the Calendar. Agents and formulas fill every other column.",
]


def arg(argv, name):
    if name in argv:
        i = argv.index(name)
        if i + 1 < len(argv):
            return argv[i + 1]
    return None


def positional(argv):
    out, skip = [], False
    for a in argv[1:]:
        if skip:
            skip = False
            continue
        if a == "--config":
            skip = True
            continue
        if a.startswith("--"):
            continue
        out.append(a)
    return out


def mappings(cfg):
    cal = cfg.get("calendar") or {}
    proof = cal.get("proof_types") or DEFAULT_PROOF_TYPES
    label = (cfg.get("swaps") or {}).get("source_label") or "long-form"
    lab = re.escape(label)
    proof_rx = "|".join(re.escape(v) for v in proof.values())
    pillars = cal.get("pillar_map") or [
        [proof_rx, "Partner & Customer Proof"],
        [r"press release|trade show", "PR & Press"],
        [r"\bpoll\b", "Engagement"],
        [r"\bpov\b", "Thought Leadership"],
        [lab + r"|newsletter|event", "Lead Generation"],
        [r"blog|video|gratitude|year-in-review", "Product & Brand Education"],
    ]
    proof_text = []
    for k, v in proof.items():
        word = {"partner": "partner news", "customer": "customer story"}.get(k.lower(), v.lower())
        proof_text.append([re.escape(v), f"LinkedIn post: {word}"])
    wtp = cal.get("what_to_post_map") or [
        [r"press release.*distribute", "Press release: wires, trade press and website (post first)"],
        [r"press release.*linkedin", "LinkedIn post: press release announcement (later the same day)"],
        *proof_text,
        [r"video", "LinkedIn video"],
        [r"blog.*publish", "Blog article on website (post first)"],
        [r"blog.*teaser", "LinkedIn post: blog teaser (2 to 4 hours after the blog)"],
        [r"trade show", "LinkedIn post: trade show"],
        [lab, f"LinkedIn post: {label} promotion"],
        [r"newsletter", "LinkedIn newsletter"],
        [r"event announcement", "LinkedIn post and LinkedIn Event listing"],
        [r"event reminder", "LinkedIn post: webinar reminder"],
        [r"live event", "Live webinar and LinkedIn post"],
        [r"\bpoll\b", "LinkedIn poll"],
        [r"\bpov\b", "LinkedIn post: opinion"],
        [r"gratitude", "LinkedIn post: thank you"],
        [r"year-in-review", "LinkedIn post: year in review"],
    ]
    return pillars, wtp


def lookup(table, ct):
    for rx, val in table:
        if re.search(rx, ct, re.I):
            return val
    return ""


def header(ws, cols, widths, fills=None):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill = PatternFill("solid", fgColor=(fills[i - 1] if fills else HEADER))
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 24


def iso(s):
    return date.fromisoformat(s) if s else None


def main(argv):
    pos = positional(argv)
    if len(pos) < 2 or "-h" in argv or "--help" in argv:
        sys.stderr.write((__doc__ or "").strip() + "\n")
        return 2
    try:
        with open(pos[0], encoding="utf-8") as f:
            data = json.load(f)
        cfg = {}
        if arg(argv, "--config"):
            with open(arg(argv, "--config"), encoding="utf-8") as f:
                cfg = json.load(f)
    except (OSError, ValueError) as e:
        sys.stderr.write(f"Cannot read input: {e}\n")
        return 2
    cal = cfg.get("calendar") or {}
    statuses = cal.get("statuses") or DEFAULT_STATUSES
    team_rx = cal.get("team_owned_regex") or DEFAULT_TEAM_OWNED
    placeholders = cal.get("proof_placeholders") or DEFAULT_PLACEHOLDERS
    week1 = (cfg.get("plan") or {}).get("week1_monday")
    pillars, wtp = mappings(cfg)
    clean, formulas = "--clean-titles" in argv, "--sheets-formulas" in argv

    rows = sorted(data.get("rows", []), key=lambda r: (r["date"], {"AM": 0, "": 1, "PM": 2}.get(r.get("ampm") or "", 1)))
    errors = []
    seq, prepared = {}, []
    for r in rows:
        if r.get("status") not in statuses:
            errors.append(f"{r.get('date')}: Status '{r.get('status')}' is not in {statuses}")
            continue
        d = date.fromisoformat(r["date"])
        wk = r.get("week")
        if wk in (None, "") and week1:
            wk = (d - date.fromisoformat(week1)).days // 7 + 1
        if wk in (None, ""):
            errors.append(f"{r['date']}: no week and no plan.week1_monday in the config")
            continue
        wk = int(wk)
        seq[wk] = seq.get(wk, 0) + 1
        ct = r["content_type"]
        title = r.get("title", "")
        if clean:
            title = re.sub(r"[\"“”]", "", title).strip()
            for ptype, ph in placeholders.items():
                if re.search(re.escape(ptype), ct, re.I) and not r.get("proof_item"):
                    title = ph
        pil, what = lookup(pillars, ct), lookup(wtp, ct)
        if not pil or not what:
            sys.stderr.write(f"Warning: no {'Pillar' if not pil else 'What to post'} mapping for '{ct}' ({r['date']}). Left blank.\n")
        notes = r.get("notes") or ("The team writes this; agents do not draft it." if re.search(team_rx, ct) else None)
        prepared.append([d, r["status"], what or None, pil or None, r.get("channel"), title, None, None,
                         r.get("visual"), notes, r.get("missing"), r.get("sources"),
                         r.get("row_id") or f"W{wk:02d}-{seq[wk]:02d}", wk, r.get("ampm") or None, ct,
                         r.get("spec"), r.get("proof_item"), None, None, None, None])
    if errors:
        sys.stderr.write("\n".join(errors) + "\n")
        return 2

    wb = Workbook()
    tw = wb.active
    tw.title = "This Week"
    ws = wb.create_sheet("Calendar")
    header(ws, COLS, WIDTHS, [HEADER] * 12 + [AGENT_HEADER] * 10)
    for p in prepared:
        ws.append(p)
    n = max(len(prepared) + 1, 2)
    last = 1000
    for i in range(2, len(prepared) + 2):
        ws.cell(row=i, column=1).number_format = "ddd d mmm yy"
        even = int(ws.cell(row=i, column=14).value) % 2 == 0
        for c, name in enumerate(COLS, 1):
            cell = ws.cell(row=i, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=name in WRAP)
            if even and c <= 12:
                cell.fill = PatternFill("solid", fgColor=BAND)
        ws.cell(row=i, column=11).font = Font(color=MISSING_TEXT)
        ws.cell(row=i, column=7).font = Font(bold=True)
    if formulas:
        ws["G2"] = G2.replace("<team>", team_rx)
        ws["H2"] = DRAFT_LINK.format(u="U2:U", v="V2:V")

    rng = f"A2:V{last}"
    for s, (bg, fg) in STATUS_FILL.items():
        if s not in statuses:
            continue
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'$B2="{s}"'], fill=PatternFill("solid", fgColor=bg),
            font=Font(color=fg, bold=s == "Approved", strike=s == "Cancelled"), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['AND($A2<>"",$A2<TODAY(),$B2<>"Posted",$B2<>"Cancelled")'],
        fill=PatternFill("solid", fgColor=OVERDUE), font=Font(bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=["$A2=TODAY()"], font=Font(bold=True)))

    for col_rng, vals in ((f"B2:B{last}", statuses), (f"O2:O{last}", ["AM", "PM"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Value", error="Choose a value from the list.")
        ws.add_data_validation(dv)
        dv.add(col_rng)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:V{n}"
    ws.column_dimensions.group("M", "V", hidden=True, outline_level=1)

    # This Week: formula view. Row 1 link, headers in row 2, formulas from row 3.
    tw.merge_cells("A1:J1")
    tw["A1"] = ("To change a Status, open the Calendar filter view \"This Week: update Status here\". "
                "Click a Draft link to open the post.")
    tw["A1"].fill = PatternFill("solid", fgColor="E3EEF7")
    tw["A1"].font = Font(bold=True)
    tw_cols = ["Date", "Status", "What to post", "Channel", "Title / angle", "Action needed", "What's missing",
               "Draft", "Visual", "Notes", "Draft URL", "Original draft"]
    for i, (h, w) in enumerate(zip(tw_cols, [14, 12, 30, 14, 44, 30, 30, 30, 30, 40, 10, 10]), 1):
        c = tw.cell(row=2, column=i, value=h)
        c.fill = PatternFill("solid", fgColor=HEADER)
        c.font = Font(bold=True, color="FFFFFF")
        tw.column_dimensions[c.column_letter].width = w
    tw.column_dimensions["K"].hidden = True
    tw.column_dimensions["L"].hidden = True
    tw.freeze_panes = "A3"
    if formulas:
        tw["A3"] = TW_A3
        tw["I3"] = TW_I3
        tw["H3"] = DRAFT_LINK.format(u="K3:K", v="L3:L")
    else:
        tw["A3"] = "Formula view: built in the live calendar (see calendar-schema.md)."

    ideas = wb.create_sheet("Ideas")
    header(ideas, ["Idea (one line)", "Added by", "Date added", "Wanted by", "Link or source", "Status",
                   "Calendar row", "Agent note"], [50, 16, 14, 14, 30, 12, 12, 50])
    for it in data.get("ideas", []):
        ideas.append([it.get("idea"), it.get("added_by"), iso(it.get("date_added")), iso(it.get("wanted_by")),
                      it.get("link"), it.get("status"), it.get("calendar_row"), it.get("agent_note")])
    dv = DataValidation(type="list", formula1='"New,Scheduled,Parked,Declined"', allow_blank=True)
    ideas.add_data_validation(dv)
    dv.add("F2:F500")
    for i in range(2, 501):
        ideas.cell(row=i, column=3).number_format = "ddd d mmm yy"
        ideas.cell(row=i, column=4).number_format = "ddd d mmm yy"
    ideas.freeze_panes = "A2"

    ms = wb.create_sheet("Milestones")
    header(ms, ["Milestone", "Type", "Start date", "End date", "Confirmed", "Notes", "Rows proposed"],
           [40, 14, 14, 14, 11, 60, 14])
    for m in data.get("milestones", []):
        ms.append([m.get("milestone"), m.get("type"), iso(m.get("start")), iso(m.get("end")),
                   m.get("confirmed"), m.get("notes"), m.get("rows_proposed")])
    for i in range(2, 201):
        ms.cell(row=i, column=3).number_format = "ddd d mmm yy"
        ms.cell(row=i, column=4).number_format = "ddd d mmm yy"
        ms.cell(row=i, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    for col_rng, vals in (("B2:B200", "Trade show,Launch,Webinar,Funding,Partner,Product"),
                          ("E2:E200", "Yes,No,Cancelled"), ("G2:G200", "Yes")):
        v = DataValidation(type="list", formula1=f'"{vals}"', allow_blank=True)
        ms.add_data_validation(v)
        v.add(col_rng)
    ms.freeze_panes = "A2"

    how = wb.create_sheet("How to use")
    header(how, ["#", "How to use this calendar"], [5, 110])
    for i, line in enumerate(HOW_TO_USE, 1):
        how.append([i, line])
        how.cell(row=i + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(pos[1])
    print(f"Wrote {pos[1]}: {len(prepared)} rows, {len(data.get('milestones', []))} milestones, "
          f"{len(data.get('ideas', []))} ideas")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
